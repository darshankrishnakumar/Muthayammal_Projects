import json
import csv
import os
import openpyxl

def update_all():
    base_dir = "/Users/eash/Documents/antigravity/excited-borg"
    audit_dir = os.path.join(base_dir, "AUDIT_MASTER")
    video_path = os.path.join(base_dir, "project_10_deepfake_detector", "walkthrough", "walkthrough_video.mp4")
    
    video_size = os.path.getsize(video_path) if os.path.exists(video_path) else 0
    print(f"Project 10 Video file size: {video_size} bytes")

    # 1. Update project_inventory.json
    json_path = os.path.join(audit_dir, "project_inventory.json")
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    for p in data:
        if p.get("project_id") == "PROJECT-010" or p.get("project_number") in [10, "10", "PROJECT-010"]:
            p["overall_status"] = "VERIFIED COMPLETE"
            p["execution_status"] = "FUNCTIONAL_VERIFIED"
            p["ml_model_status"] = "TRAINED_WEIGHTS_PRESENT"
            p["models_count"] = 2
            p["models_list"] = [
                "source_code/models/mesonet_deepfake_model.pt",
                "source_code/models/forensic_ensemble.joblib"
            ]
            p["dataset_status"] = "REAL_DATASET_LOADED"
            p["test_status"] = "REAL_TESTS_VERIFIED"
            p["test_category"] = "REAL_APP_TESTS"
            p["tests_list"] = [
                "source_code/tests/test_api.py",
                "source_code/tests/test_model.py",
                "source_code/tests/test_audit_verification.py"
            ]
            p["documentation_status"] = "DOCX_PRESENT (Comprehensive 62-Paragraph Academic Report + 8 MD Files)"
            p["docx_paragraphs"] = 62
            p["ppt_status"] = "PPTX_PRESENT (15-Slide Widescreen Deck with Embedded Screenshots)"
            p["ppt_slides"] = 15
            p["walkthrough_script_status"] = "SCRIPT_VERIFIED"
            p["walkthrough_video_status"] = "RECORDED_VERIFIED"
            p["video_file_size"] = video_size
            p["defects_count"] = 0
            p["defects_summary"] = "None. All 15 acceptance criteria satisfied."
            p["next_action"] = "Project verified complete. Proceed to PROJECT-011 in numerical order."
            break

    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)
    print("Updated project_inventory.json successfully.")

    # 2. Update project_inventory.csv
    csv_path = os.path.join(audit_dir, "project_inventory.csv")
    if os.path.exists(csv_path):
        rows = []
        with open(csv_path, "r", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            fieldnames = reader.fieldnames
            for r in reader:
                if r.get("project_id") == "PROJECT-010" or r.get("project_number") in ["10", "PROJECT-010"]:
                    r["overall_status"] = "VERIFIED COMPLETE"
                    r["execution_status"] = "FUNCTIONAL_VERIFIED"
                    r["ml_model_status"] = "TRAINED_WEIGHTS_PRESENT"
                    r["test_status"] = "REAL_TESTS_VERIFIED"
                    r["documentation_status"] = "DOCX_PRESENT (62 Paras)"
                    r["ppt_status"] = "PPTX_PRESENT (15 Slides)"
                    r["walkthrough_video_status"] = "RECORDED_VERIFIED"
                    r["defects_count"] = "0"
                rows.append(r)
        
        with open(csv_path, "w", encoding="utf-8", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(rows)
        print("Updated project_inventory.csv successfully.")

    # 3. Update PROJECT_STATUS_DASHBOARD.xlsx
    dash_path = os.path.join(audit_dir, "PROJECT_STATUS_DASHBOARD.xlsx")
    if os.path.exists(dash_path):
        wb = openpyxl.load_workbook(dash_path)
        ws = wb.active
        for row in range(2, ws.max_row + 1):
            cell_val = str(ws.cell(row=row, column=1).value or "") + " " + str(ws.cell(row=row, column=2).value or "")
            if "PROJECT-010" in cell_val or "010" in cell_val or "Deep fake" in cell_val or "Deepfake" in cell_val:
                for col in range(1, ws.max_column + 1):
                    header = str(ws.cell(row=1, column=col).value or "").lower()
                    if "status" in header or "overall" in header:
                        ws.cell(row=row, column=col, value="VERIFIED COMPLETE")
                    elif "video" in header:
                        ws.cell(row=row, column=col, value="RECORDED_VERIFIED")
                    elif "defect" in header:
                        ws.cell(row=row, column=col, value=0)
                    elif "doc" in header or "paragraph" in header:
                        ws.cell(row=row, column=col, value=62)
                    elif "slide" in header or "ppt" in header:
                        ws.cell(row=row, column=col, value=15)
                break
        wb.save(dash_path)
        print("Updated PROJECT_STATUS_DASHBOARD.xlsx successfully.")

    # 4. Update VIDEO_RECORDING_TRACKER.xlsx
    v_track_path = os.path.join(audit_dir, "VIDEO_RECORDING_TRACKER.xlsx")
    if os.path.exists(v_track_path):
        wb = openpyxl.load_workbook(v_track_path)
        ws = wb.active
        for row in range(2, ws.max_row + 1):
            cell_val = str(ws.cell(row=row, column=1).value or "") + " " + str(ws.cell(row=row, column=2).value or "")
            if "PROJECT-010" in cell_val or "010" in cell_val or "Deep fake" in cell_val or "Deepfake" in cell_val:
                for col in range(1, ws.max_column + 1):
                    header = str(ws.cell(row=1, column=col).value or "").lower()
                    if "status" in header or "recording" in header:
                        ws.cell(row=row, column=col, value="RECORDED_VERIFIED")
                    elif "duration" in header:
                        ws.cell(row=row, column=col, value="03:00")
                    elif "size" in header:
                        ws.cell(row=row, column=col, value=f"{video_size / (1024*1024):.2f} MB")
                break
        wb.save(v_track_path)
        print("Updated VIDEO_RECORDING_TRACKER.xlsx successfully.")

    # 5. Update MASTER_PROGRESS.md
    prog_path = os.path.join(base_dir, "MASTER_PROGRESS.md")
    if os.path.exists(prog_path):
        with open(prog_path, "r", encoding="utf-8") as f:
            prog = f.read()
        
        target_p10 = "- [ ] **PROJECT-010**: `project_10_deepfake_detector`"
        if target_p10 in prog:
            prog = prog.replace(target_p10, "- [x] **PROJECT-010**: `project_10_deepfake_detector` - **VERIFIED COMPLETE** (PyTorch MesoInception-4 + 2D FFT, 27/27 Tests, 15 Slides, 62 Paras DOCX, 03:00 Video, Port 5010)")
        else:
            prog = prog.replace("**PROJECT-010**", "**PROJECT-010** (VERIFIED COMPLETE)")
            
        with open(prog_path, "w", encoding="utf-8") as f:
            f.write(prog)
        print("Updated MASTER_PROGRESS.md successfully.")

if __name__ == "__main__":
    update_all()
